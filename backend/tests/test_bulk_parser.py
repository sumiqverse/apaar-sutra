"""Tests for bulk file template generation and streaming parsers (CSV and XLSX)."""

from datetime import date, datetime
import io
import openpyxl
import pytest

from app.services.bulk_ingestion_service import BulkIngestionService


def test_generate_csv_template():
    """CSV template contains valid header and example rows."""
    content, media_type = BulkIngestionService.generate_template("csv")
    assert media_type == "text/csv"
    text = content.decode("utf-8-sig")
    lines = text.strip().split("\n")
    assert len(lines) >= 2
    header = lines[0]
    assert "apaar_id" in header
    assert "student_ref_id" in header
    assert "quota_type" in header


def test_generate_xlsx_template():
    """XLSX template contains 3 sheets, styling, and data dictionary."""
    content, media_type = BulkIngestionService.generate_template("xlsx")
    assert media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(content))
    sheet_names = wb.sheetnames
    assert "Student Enrolment" in sheet_names
    assert "Data Dictionary" in sheet_names
    assert "Program Quotas" in sheet_names

    ws = wb["Student Enrolment"]
    headers = [cell.value for cell in ws[1]]
    assert "apaar_id" in headers
    assert "student_ref_id" in headers
    assert "admission_category" in headers


def test_parse_csv_with_utf8_bom(tmp_path):
    """CSV parser transparently handles UTF-8 BOM."""
    csv_file = tmp_path / "students_bom.csv"
    bom_content = (
        "\ufeffapaar_id,student_ref_id,name,dob,gender,program_name,branch,academic_year\n"
        "200000000010,REF-001,Amit Kumar,2003-01-15,MALE,BTech CS,CSE,2025-26\n"
        "200000000020,REF-002,Pooja Sharma,2004-02-20,FEMALE,BTech IT,IT,2025-26\n"
    )
    csv_file.write_text(bom_content, encoding="utf-8-sig")

    batches = list(BulkIngestionService.stream_parse_file(str(csv_file), chunk_size=10))
    assert len(batches) == 1
    rows = batches[0]
    assert len(rows) == 2
    row_idx_1, rec_1 = rows[0]
    row_idx_2, rec_2 = rows[1]
    assert rec_1["apaar_id"] == "200000000010"
    assert rec_1["name"] == "Amit Kumar"
    assert rec_2["apaar_id"] == "200000000020"


def test_parse_csv_semicolon_delimiter_and_whitespace(tmp_path):
    """CSV parser auto-detects semicolon delimiter and strips whitespace."""
    csv_file = tmp_path / "students_semi.csv"
    content = (
        "apaar_id;student_ref_id;name;dob;gender;program_name;branch;academic_year\n"
        "  200000000030  ; REF-003 ; Sunita Rao ; 2003-05-10 ; FEMALE ; BTech ; EE ; 2025-26 \n"
        "\n"  # blank line
        "  200000000040  ; REF-004 ; Rajesh Nair ; 2002-12-01 ; MALE ; BTech ; ME ; 2025-26 \n"
    )
    csv_file.write_text(content, encoding="utf-8")

    batches = list(BulkIngestionService.stream_parse_file(str(csv_file), chunk_size=10))
    assert len(batches) == 1
    rows = batches[0]
    assert len(rows) == 2
    _, rec_1 = rows[0]
    _, rec_2 = rows[1]
    assert rec_1["apaar_id"] == "200000000030"
    assert rec_1["name"] == "Sunita Rao"
    assert rec_2["apaar_id"] == "200000000040"


def test_parse_csv_hindi_unicode_support(tmp_path):
    """CSV parser preserves Hindi Devanagari text and numerals."""
    csv_file = tmp_path / "students_hindi.csv"
    content = (
        "apaar_id,student_ref_id,name,dob,gender,program_name,branch,academic_year\n"
        "२०००००००००५०,REF-005,रोहित शर्मा,2003-08-15,MALE,BTech CS,CSE,2025-26\n"
    )
    csv_file.write_text(content, encoding="utf-8")

    batches = list(BulkIngestionService.stream_parse_file(str(csv_file), chunk_size=10))
    assert len(batches) == 1
    rows = batches[0]
    assert len(rows) == 1
    _, rec = rows[0]
    assert rec["apaar_id"] == "२०००००००००५०"
    assert rec["name"] == "रोहित शर्मा"


def test_parse_xlsx_with_excel_date_serials(tmp_path):
    """XLSX parser converts Excel numeric date serials to ISO format strings."""
    xlsx_file = tmp_path / "students_serial.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Enrolment"
    ws.append(["apaar_id", "student_ref_id", "name", "dob", "gender", "academic_year"])
    # 38122 is Excel serial for 2004-05-15
    ws.append(["200000000060", "REF-006", "Kavita Reddy", 38122, "FEMALE", "2025-26"])
    wb.save(xlsx_file)

    batches = list(BulkIngestionService.stream_parse_file(str(xlsx_file), chunk_size=10))
    assert len(batches) == 1
    rows = batches[0]
    assert len(rows) == 1
    _, rec = rows[0]
    assert rec["apaar_id"] == "200000000060"
    assert rec["dob"] == "2004-05-15"
